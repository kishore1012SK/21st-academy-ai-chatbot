"""
Fully local RAG pipeline:
  extract text -> chunk -> embed via Ollama (nomic-embed-text) -> store as
  .npy files under vector_db/ -> cosine-similarity search at query time.

No data ever leaves the machine running Ollama + this backend.
"""
import csv
import uuid
import numpy as np
from pathlib import Path
from typing import List, Tuple

import pypdf
import docx
import openpyxl

from config import settings
from services.ollama_service import embed_text


def extract_text(filepath: Path) -> str:
    ext = filepath.suffix.lower()

    if ext == ".pdf":
        text = []
        reader = pypdf.PdfReader(str(filepath))
        for page in reader.pages:
            text.append(page.extract_text() or "")
        return "\n".join(text)

    if ext == ".docx":
        d = docx.Document(str(filepath))
        return "\n".join(p.text for p in d.paragraphs)

    if ext == ".txt":
        return filepath.read_text(errors="ignore")

    if ext == ".csv":
        rows = []
        with open(filepath, newline="", errors="ignore") as f:
            for row in csv.reader(f):
                rows.append(", ".join(row))
        return "\n".join(rows)

    if ext in (".xlsx", ".xls"):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                lines.append(", ".join(str(c) for c in row if c is not None))
        return "\n".join(lines)

    if ext in (".pptx", ".ppt"):
        import pptx
        prs = pptx.Presentation(str(filepath))
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text.append(shape.text)
        return "\n".join(text)

    raise ValueError(f"Unsupported file type: {ext}")


def chunk_text(text: str, size: int = None, overlap: int = None) -> List[str]:
    size = size or settings.CHUNK_SIZE
    overlap = overlap or settings.CHUNK_OVERLAP
    text = " ".join(text.split())  # normalize whitespace
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunks.append(text[start:end])
        start = end - overlap
        if start < 0 or end >= len(text):
            break
    return [c.strip() for c in chunks if c.strip()]


async def embed_and_store_chunks(document_id: str, chunks: List[str]) -> List[Tuple[str, str]]:
    """Embeds each chunk and stores the vector as .npy. Returns [(chunk_text, embedding_path)]."""
    results = []
    doc_dir = settings.VECTOR_DB_DIR / document_id
    doc_dir.mkdir(parents=True, exist_ok=True)

    for i, chunk in enumerate(chunks):
        vector = await embed_text(chunk)
        vec_path = doc_dir / f"chunk_{i}_{uuid.uuid4().hex[:8]}.npy"
        np.save(vec_path, np.array(vector, dtype=np.float32))
        results.append((chunk, str(vec_path)))

    return results


def _cosine_sim(a: np.ndarray, b: np.ndarray) -> float:
    denom = (np.linalg.norm(a) * np.linalg.norm(b))
    if denom == 0:
        return 0.0
    return float(np.dot(a, b) / denom)


async def search_similar_chunks(query: str, db_session, top_k: int = None) -> List[str]:
    """Embeds the query and does a brute-force cosine search over stored chunk vectors.
    Fine for a knowledge base of a few thousand chunks; swap for FAISS/Chroma if it
    ever needs to scale further — the interface stays the same either way."""
    from database.models import DocumentChunk  # local import avoids circulars

    top_k = top_k or settings.RAG_TOP_K
    query_vec = np.array(await embed_text(query), dtype=np.float32)
    if query_vec.size == 0:
        return []

    all_chunks = db_session.query(DocumentChunk).filter(DocumentChunk.embedding_path.isnot(None)).all()
    scored = []
    for chunk in all_chunks:
        try:
            vec = np.load(chunk.embedding_path)
        except FileNotFoundError:
            continue
        score = _cosine_sim(query_vec, vec)
        scored.append((score, chunk.content))

    scored.sort(key=lambda x: x[0], reverse=True)
    return [content for score, content in scored[:top_k] if score > 0.3]
